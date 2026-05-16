from io import BytesIO
from html import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


def is_arabic_text(text):
    """
    Simple Arabic character detection.
    """
    if not text:
        return False

    for char in text:
        if "\u0600" <= char <= "\u06FF":
            return True

    return False


def create_word_export(brief):
    """
    Create a Word .docx file in memory.
    """
    document = Document()

    language = brief["language"]
    is_arabic = language == "ar" or is_arabic_text(brief["brief"])

    title = "LEAF AI Designer - Saved Design Brief"
    if is_arabic:
        title = "مصمم LEAF الذكي - فكرة تصميم محفوظة"

    document.add_heading(title, level=1)

    metadata = [
        ("ID", brief["id"]),
        ("Created At", brief["created_at"]),
        ("Language", brief["language"]),
        ("Prompt", brief["prompt"]),
        ("References", brief["reference_names"] or "Not specified"),
    ]

    if is_arabic:
        metadata = [
            ("المعرّف", brief["id"]),
            ("تاريخ الإنشاء", brief["created_at"]),
            ("اللغة", brief["language"]),
            ("الوصف", brief["prompt"]),
            ("المراجع", brief["reference_names"] or "غير محدد"),
        ]

    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value))

    document.add_paragraph("")

    for line in brief["brief"].splitlines():
        clean_line = line.strip()

        if not clean_line:
            document.add_paragraph("")
            continue

        if clean_line.startswith("#"):
            heading_text = clean_line.replace("#", "").strip()
            level = min(clean_line.count("#"), 3)
            document.add_heading(heading_text, level=level)
        elif clean_line.startswith("- "):
            document.add_paragraph(clean_line[2:], style="List Bullet")
        elif clean_line.startswith(">"):
            document.add_paragraph(clean_line.replace(">", "").strip())
        else:
            paragraph = document.add_paragraph(clean_line)
            for run in paragraph.runs:
                run.font.size = Pt(11)

    file_buffer = BytesIO()
    document.save(file_buffer)
    file_buffer.seek(0)

    return file_buffer.getvalue()


def clean_markdown_for_pdf(text):
    """
    Convert simple markdown into PDF-friendly text.
    """
    lines = []

    for line in text.splitlines():
        clean_line = line.strip()

        if clean_line.startswith("#"):
            clean_line = clean_line.replace("#", "").strip()

        if clean_line.startswith("- "):
            clean_line = "• " + clean_line[2:]

        if clean_line.startswith(">"):
            clean_line = clean_line.replace(">", "").strip()

        clean_line = clean_line.replace("**", "")

        lines.append(clean_line)

    return "\n".join(lines)


def create_pdf_export(brief):
    """
    Create a PDF file in memory.
    """
    file_buffer = BytesIO()

    doc = SimpleDocTemplate(
        file_buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    is_arabic = brief["language"] == "ar" or is_arabic_text(brief["brief"])

    alignment = TA_RIGHT if is_arabic else TA_LEFT

    title_style = ParagraphStyle(
        "LeafTitle",
        parent=styles["Title"],
        alignment=alignment,
        fontSize=18,
        leading=24,
        spaceAfter=16
    )

    normal_style = ParagraphStyle(
        "LeafNormal",
        parent=styles["Normal"],
        alignment=alignment,
        fontSize=10,
        leading=15,
        spaceAfter=8
    )

    title = "LEAF AI Designer - Saved Design Brief"
    if is_arabic:
        title = "مصمم LEAF الذكي - فكرة تصميم محفوظة"

    story = []
    story.append(Paragraph(escape(title), title_style))
    story.append(Spacer(1, 12))

    metadata_lines = [
        f"ID: {brief['id']}",
        f"Created At: {brief['created_at']}",
        f"Language: {brief['language']}",
        f"Prompt: {brief['prompt']}",
        f"References: {brief['reference_names'] or 'Not specified'}",
    ]

    if is_arabic:
        metadata_lines = [
            f"المعرّف: {brief['id']}",
            f"تاريخ الإنشاء: {brief['created_at']}",
            f"اللغة: {brief['language']}",
            f"الوصف: {brief['prompt']}",
            f"المراجع: {brief['reference_names'] or 'غير محدد'}",
        ]

    for line in metadata_lines:
        story.append(Paragraph(escape(str(line)), normal_style))

    story.append(Spacer(1, 12))

    clean_brief = clean_markdown_for_pdf(brief["brief"])

    for line in clean_brief.splitlines():
        if line.strip():
            story.append(Paragraph(escape(line.strip()), normal_style))
        else:
            story.append(Spacer(1, 8))

    doc.build(story)

    file_buffer.seek(0)
    return file_buffer.getvalue()